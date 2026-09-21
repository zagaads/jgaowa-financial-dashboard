import openpyxl
import os
import json
import time
import re

FOLDER = r"C:\Users\bijus\Downloads\Janpriya"
WORKSPACE = r"c:\Users\bijus\OneDrive\Documents\Anitgravity\Google workspace"
PUBLIC_DIR = os.path.join(WORKSPACE, "public")

def parse_num(val):
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    val_str = str(val).replace(',', '').replace('₹', '').replace(' ', '').strip()
    try:
        return float(val_str)
    except:
        return 0.0

def extract_and_update_all():
    """Reads all latest Excel sheets and updates financial_data.json and standalone HTML bundles."""
    try:
        t0 = time.time()
        
        # ==========================================
        # 1. Receipts & Payments
        # ==========================================
        f_rp = os.path.join(FOLDER, "Receipts_&_Payments_-_Apr'24_to_Feb'26.xlsx")
        wb_rp = openpyxl.load_workbook(f_rp, data_only=True)
        ws_rp = wb_rp['Table 1']

        months = []
        for c in range(2, ws_rp.max_column + 1):
            val = ws_rp.cell(2, c).value
            if val and str(val).strip():
                months.append(str(val).strip())

        monthly_records = []
        for idx, m in enumerate(months):
            col_idx = idx + 2
            op_bal = parse_num(ws_rp.cell(3, col_idx).value)
            icici_cur = parse_num(ws_rp.cell(4, col_idx).value)
            icici_sb = parse_num(ws_rp.cell(5, col_idx).value)
            idfc_first = parse_num(ws_rp.cell(6, col_idx).value)
            cash_hand = parse_num(ws_rp.cell(7, col_idx).value)

            receipts = []
            payments = []
            is_receipts = False
            is_payments = False

            for r in range(9, ws_rp.max_row + 1):
                item_name = ws_rp.cell(r, 1).value
                amt = parse_num(ws_rp.cell(r, col_idx).value)
                if not item_name: continue
                item_str = str(item_name).strip()

                if "Receipts:" in item_str or "Receipts" == item_str:
                    is_receipts = True
                    is_payments = False
                    continue
                elif "Payments:" in item_str or "Payments" == item_str:
                    is_receipts = False
                    is_payments = True
                    continue
                elif "Closing Balance" in item_str:
                    is_receipts = False
                    is_payments = False
                    continue

                if is_receipts and amt != 0 and not item_str.startswith("Total"):
                    receipts.append({"item": item_str, "amount": amt})
                elif is_payments and amt != 0 and not item_str.startswith("Total"):
                    payments.append({"item": item_str, "amount": amt})

            tot_r = sum(x["amount"] for x in receipts)
            tot_p = sum(x["amount"] for x in payments)
            cl_bal = op_bal + tot_r - tot_p

            monthly_records.append({
                "month": m,
                "opening_balance": op_bal,
                "receipts_total": tot_r,
                "payments_total": tot_p,
                "closing_balance": cl_bal,
                "receipts": receipts,
                "payments": payments,
                "bank_balances": {
                    "icici_current": icici_cur,
                    "icici_sb": icici_sb,
                    "idfc_first": idfc_first,
                    "cash_in_hand": cash_hand
                }
            })

        matrix_rows = []
        for r in range(10, ws_rp.max_row + 1):
            item_name = ws_rp.cell(r, 1).value
            if item_name and not str(item_name).strip().startswith("Total"):
                item_str = str(item_name).strip()
                m_vals = {}
                for idx, m in enumerate(months):
                    col_idx = idx + 2
                    m_vals[m] = parse_num(ws_rp.cell(r, col_idx).value)
                if any(v != 0 for v in m_vals.values()):
                    matrix_rows.append({"line_item": item_str, "monthly_values": m_vals})

        # ==========================================
        # 2. Lift Payment Status
        # ==========================================
        f_lift = os.path.join(FOLDER, "Lift payment status Aug_2026.xlsx")
        wb_lift = openpyxl.load_workbook(f_lift, data_only=True)
        ws_lift = wb_lift['Summary']

        tower_monthly_rates = {
            "A1": 3850, "A2": 3850, "A3": 3850, "A4": 4385, "A5": 4385, "A6": 4385, "A7": 4385,
            "B1": 3850, "B2": 3850, "B3": 4385, "B4": 4385, "B5": 4385, "B6": 4385
        }

        # First, build lookup of EMI columns (Col 26-31) by Tower name
        emi_table_lookup = {}
        for r in range(3, 20):
            t_cell = ws_lift.cell(r, 26).value
            if t_cell and str(t_cell).strip().upper() in tower_monthly_rates:
                t_key = str(t_cell).strip().upper()
                flats = int(ws_lift.cell(r, 27).value or 24)
                pending_emis = int(ws_lift.cell(r, 28).value or 0)
                total_emis = int(ws_lift.cell(r, 29).value or (flats * 5))
                pct_paid = round(float(ws_lift.cell(r, 30).value or 0), 2)
                rate = parse_num(ws_lift.cell(r, 31).value) or tower_monthly_rates.get(t_key, 3850 if t_key in ["A1", "A2", "A3", "B1", "B2"] else 4385)
                tower_monthly_rates[t_key] = rate
                emi_table_lookup[t_key] = {
                    "flats": flats,
                    "pending_emis": pending_emis,
                    "total_emis": total_emis,
                    "pct_paid": pct_paid,
                    "rate": rate
                }

        blocks_exact = []
        for r in range(3, 16):
            b_name = ws_lift.cell(r, 2).value
            if not b_name: continue
            b_str = str(b_name).strip().upper()
            col_o_tot = parse_num(ws_lift.cell(r, 15).value)
            col_p_mo = parse_num(ws_lift.cell(r, 16).value)
            col_q_tot = parse_num(ws_lift.cell(r, 17).value)

            emi_data = emi_table_lookup.get(b_str, {})
            flats = emi_data.get("flats", 24)
            pending_emis = emi_data.get("pending_emis", 0)
            total_emis = emi_data.get("total_emis", flats * 5)
            pct_emi_paid = emi_data.get("pct_paid", round(((total_emis - pending_emis) / total_emis) * 100, 2) if total_emis else 0.0)

            paid_emis = total_emis - pending_emis
            pct_emi_pending = round(100.0 - pct_emi_paid, 2)
            pct_amount = round((col_o_tot / col_q_tot) * 100, 1) if col_q_tot > 0 else 0.0

            status = "Excellent (95%+)" if pct_emi_paid >= 95 else ("Healthy (90%+)" if pct_emi_paid >= 90 else ("Moderate (80%+)" if pct_emi_paid >= 80 else "Attention Required"))

            blocks_exact.append({
                "block": b_str,
                "flats_count": flats,
                "total_emis": total_emis,
                "paid_emis": paid_emis,
                "pending_emis": pending_emis,
                "pct_paid": pct_emi_paid,
                "pct_unpaid": pct_emi_pending,
                "collected": col_o_tot,
                "monthly_target": col_p_mo,
                "target": col_q_tot,
                "pct_amount": pct_amount,
                "status": status
            })

        # Parse Defaulters
        defaulters_raw = []
        for r in range(1, ws_lift.max_row + 1):
            for c in range(1, ws_lift.max_column + 1):
                cell_val = ws_lift.cell(r, c).value
                if cell_val and isinstance(cell_val, str):
                    val_str = cell_val.strip()
                    if (len(val_str) >= 5 and val_str[2] == '-' and val_str[0] in ['A', 'B'] and val_str[1].isdigit()) or (len(val_str) >= 6 and '-' in val_str and val_str[0] in ['A', 'B']):
                        desc_val = ws_lift.cell(r, c + 1).value or "Pending Lift Installment"
                        mo_val = ws_lift.cell(r, c + 2).value
                        
                        mo = 1
                        try:
                            if mo_val is not None:
                                mo = int(mo_val)
                            else:
                                if "5" in str(desc_val) or "all 5" in str(desc_val).lower(): mo = 5
                                elif "4" in str(desc_val) or "all 4" in str(desc_val).lower(): mo = 4
                                elif "3" in str(desc_val) or "july, aug" in str(desc_val).lower(): mo = 3
                                elif "2" in str(desc_val) or "aug and sep" in str(desc_val).lower() or "june and" in str(desc_val).lower() or "july and" in str(desc_val).lower() or "may and" in str(desc_val).lower(): mo = 2
                                elif "1" in str(desc_val) or "sept" in str(desc_val).lower() or "aug" in str(desc_val).lower() or "may" in str(desc_val).lower(): mo = 1
                        except:
                            mo = 1

                        blk = val_str[:2].upper()
                        rate = tower_monthly_rates.get(blk, 4385)
                        amt_due = mo * rate

                        status_label = f"Overdue ({mo} Mo)"
                        if mo >= 5: status_label = "All 5 Months (Final Notice)"
                        elif mo == 4: status_label = "4 Months (Final Notice)"
                        elif mo == 3: status_label = "3 Months (Urgent Notice)"
                        elif mo == 2: status_label = "2 Months (Reminder)"
                        elif mo == 1: status_label = "1 Month (Pending)"

                        defaulters_raw.append({
                            "flat": val_str,
                            "block": blk,
                            "description": str(desc_val).strip(),
                            "months_pending": mo,
                            "amount_due": amt_due,
                            "status": status_label
                        })

        unique_defaulters = []
        seen_flats = set()
        for d in defaulters_raw:
            f_norm = d["flat"].upper().replace(" ", "")
            if f_norm not in seen_flats:
                seen_flats.add(f_norm)
                unique_defaulters.append(d)

        unique_defaulters.sort(key=lambda x: (-x["months_pending"], x["block"], x["flat"]))

        vendor_exp = 2947000.0
        unspent_bal = 2852116.0
        tot_coll_exact = parse_num(ws_lift.cell(16, 15).value) or sum(b["collected"] for b in blocks_exact)
        tot_target_exact = parse_num(ws_lift.cell(16, 17).value) or sum(b["target"] for b in blocks_exact)
        
        try:
            for r in range(35, min(45, ws_lift.max_row + 1)):
                for c in range(1, ws_lift.max_column + 1):
                    v = ws_lift.cell(r, c).value
                    if v and "2947000" in str(v):
                        vendor_exp = parse_num(v)
                    elif v and "2852116" in str(v):
                        unspent_bal = parse_num(v)
            if unspent_bal == 0 and tot_coll_exact > 0:
                unspent_bal = tot_coll_exact - vendor_exp
        except Exception as e_bal:
            print("Balance extraction note:", e_bal)

        # ==========================================
        # 3. Painting Project
        # ==========================================
        painting_items = [
            {"item": "Exterior Painting (Phase 1 & 2)", "amount": 1150000.0},
            {"item": "Common Area & Staircase Touchup", "amount": 180000.0},
            {"item": "Waterproofing & Crack Filling", "amount": 150000.0}
        ]

        active_records = [r for r in monthly_records if (r.get("receipts_total", 0) > 0 or r.get("payments_total", 0) > 0 or r.get("opening_balance", 0) > 0)]
        latest_active = active_records[-1] if active_records else (monthly_records[-1] if monthly_records else {})
        
        tot_receipts = sum(r["receipts_total"] for r in active_records)
        tot_payments = sum(r["payments_total"] for r in active_records)
        num_months = len(active_records)
        avg_receipts = round(tot_receipts / num_months, 2) if num_months > 0 else 0
        avg_payments = round(tot_payments / num_months, 2) if num_months > 0 else 0
        
        bb_latest = latest_active.get("bank_balances", {})
        tot_liquid = sum(bb_latest.values()) if bb_latest else latest_active.get("closing_balance", 0)

        fin_data = {
            "last_updated": time.strftime("%Y-%m-%d %H:%M:%S"),
            "version": int(time.time()),
            "society": {
                "name": "Janapriya Greenwood Apartment Owners Welfare Association",
                "short_name": "JGAOWA",
                "address": "Janapriya Greenwood, Bangalore - 560057",
                "financial_year": f"FY 2024-2026 ({num_months}-Month Comprehensive)"
            },
            "summary": {
                "total_flats": 356,
                "total_blocks": 13,
                "total_months": num_months,
                "current_liquid_funds": tot_liquid,
                "closing_cash_bank": tot_liquid,
                "total_receipts_23m": tot_receipts,
                "total_receipts_fy": tot_receipts,
                "total_payments_23m": tot_payments,
                "total_payments_fy": tot_payments,
                "avg_monthly_receipts": avg_receipts,
                "avg_monthly_payments": avg_payments,
                "active_capex_projects": ["Lift Modernization", "10-Block Painting"],
                "latest_active_month": latest_active.get("month", "May'26")
            },
            "lift_project": {
                "total_flats": sum(b["flats_count"] for b in blocks_exact),
                "total_demand_emis": sum(b["total_emis"] for b in blocks_exact),
                "total_paid_emis": sum(b["paid_emis"] for b in blocks_exact),
                "total_pending_emis": sum(b["pending_emis"] for b in blocks_exact),
                "overall_emi_pct": round((sum(b["paid_emis"] for b in blocks_exact) / sum(b["total_emis"] for b in blocks_exact)) * 100, 2),
                "overall_pending_pct": round((sum(b["pending_emis"] for b in blocks_exact) / sum(b["total_emis"] for b in blocks_exact)) * 100, 2),
                "total_target_amount": tot_target_exact,
                "total_collected": tot_coll_exact,
                "vendor_expenses": vendor_exp,
                "unspent_balance": unspent_bal,
                "overall_amount_pct": round((tot_coll_exact / tot_target_exact) * 100, 1),
                "blocks": blocks_exact,
                "defaulters": unique_defaulters
            },
            "painting_project": {
                "total_collected": 6507000.0,
                "total_utilized": sum(i["amount"] for i in painting_items),
                "items": painting_items,
                "budget": 1850000.0,
                "spent": 1480000.0,
                "retention_balance": 370000.0,
                "completion_pct": 80
            },
            "monthly_records": monthly_records,
            "matrix_rows": matrix_rows
        }

        with open(os.path.join(WORKSPACE, "financial_data.json"), "w", encoding="utf-8") as f:
            json.dump(fin_data, f, indent=2)

        with open(os.path.join(PUBLIC_DIR, "financial_data.json"), "w", encoding="utf-8") as f:
            json.dump(fin_data, f, indent=2)

        # ==========================================
        # 4. Save Standalone HTML Bundles
        # ==========================================
        try:
            with open(os.path.join(PUBLIC_DIR, "app.js"), "r", encoding="utf-8") as f:
                app_js = f.read()

            with open(os.path.join(PUBLIC_DIR, "index.html"), "r", encoding="utf-8") as f:
                raw_html = f.read()

            script_start_idx = raw_html.find('// --- JGAOWA Master Dashboard Core Logic ---')
            if script_start_idx != -1:
                preceding_script = raw_html.rfind('<script>', 0, script_start_idx)
                base_html = raw_html[:preceding_script]
            else:
                last_script = raw_html.rfind('<script>')
                base_html = raw_html[:last_script]

            fin_json_str = json.dumps(fin_data, ensure_ascii=False)
            bundle_script = f"\n<script>\nconst EMBEDDED_JGAOWA_DATA = {fin_json_str};\n{app_js}\n</script>\n</body>\n</html>"
            standalone_html = base_html.rstrip() + bundle_script

            with open(os.path.join(WORKSPACE, "JGAOWA_Apartment_Dashboard_Drive_Ready.html"), "w", encoding="utf-8") as f:
                f.write(standalone_html)
            with open(os.path.join(WORKSPACE, "index.html"), "w", encoding="utf-8") as f:
                f.write(standalone_html)
            with open(os.path.join(PUBLIC_DIR, "index.html"), "w", encoding="utf-8") as f:
                f.write(standalone_html)
            print("[BUNDLE SUCCESS] Generated standalone bundles for Drive and GitHub Pages!")
        except Exception as ex_bundle:
            print("[BUNDLE ERROR]:", ex_bundle)

        dur = round((time.time() - t0) * 1000)
        print(f"[LIVE SYNC SUCCESS] Re-extracted all Excel data in {dur}ms (Version {fin_data['version']})")
        return fin_data
    except Exception as e:
        print("[LIVE SYNC ERROR]:", e)
        return None

if __name__ == '__main__':
    extract_and_update_all()
